'''
对问题反馈人(Top5)反馈的问题评级占比分析，输出柱形图 ------ 江坚
'''

import openpyxl
from openpyxl.chart import BarChart, Series, Reference


def data_create(dateResource):
    dadaist = []
    jxx_dict = {}
    jxx_dicta = {}
    jxx_dictb = {}
    jxx_dictc = {}
    jxx_dictd = {}
    for j in range(0, len(dateResource)):
        str4 = dateResource[j][3]
        str1 = str4[:str4.index('\n')]
        stra = dateResource[j][8]
        if stra == "A":
            if str1 in jxx_dicta:
                jxx_dicta[str1] = (jxx_dicta.get(str1) + 1)
            else:
                jxx_dicta.setdefault(str1, 1)
        elif stra == "B":
            if str1 in jxx_dictb:
                jxx_dictb[str1] = (jxx_dictb.get(str1) + 1)
            else:
                jxx_dictb.setdefault(str1, 1)
        elif stra == "C":
            if str1 in jxx_dictc:
                jxx_dictc[str1] = (jxx_dictc.get(str1) + 1)
            else:
                jxx_dictc.setdefault(str1, 1)
        else:
            if str1 in jxx_dictd:
                jxx_dictd[str1] = (jxx_dictd.get(str1) + 1)
            else:
                jxx_dictd.setdefault(str1, 1)
        if str1 in jxx_dict:
            jxx_dict[str1] = (jxx_dict.get(str1) + 1)
        else:
            jxx_dict.setdefault(str1, 1)
    a1 = sorted(jxx_dict.items(), key=lambda x: x[1], reverse=True)
    if len(a1) > 5:
        b1 = 5
    else:
        b1 = len(a1)
    list0 = ["姓名", "A", "B", "C", "D"]
    dadaist.append(list0)
    for m in range(0, b1):
        listrange = [a1[m][0], jxx_dicta.get(a1[m][0]), jxx_dictb.get(a1[m][0]), jxx_dictc.get(a1[m][0]),
                     jxx_dictd.get(a1[m][0])]
        dadaist.append(listrange)
    return dadaist


def draw_top5(dadaist):
    wb = openpyxl.load_workbook("total.xlsx")
    sheets = wb.get_sheet_names()
    ws = wb.create_sheet(title="top5", index=len(sheets))
    rows = dadaist

    for row in rows:
        ws.append(row)

    chart1 = BarChart()
    chart1.type = "col"
    chart1.style = 10
    chart1.title = "Top5反馈评级数"
    chart1.y_axis.title = '数量'
    chart1.x_axis.title = '人名'

    data = Reference(ws, min_col=2, min_row=1, max_row=len(rows), max_col=len(rows[0]))
    cats = Reference(ws, min_col=1, min_row=2, max_row=len(rows))
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    chart1.shape = 4
    ws.add_chart(chart1, "A10")

    wb.save("total.xlsx")
